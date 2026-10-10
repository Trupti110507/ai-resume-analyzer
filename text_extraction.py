"""
Resume Upload Module (file checks) and Text Extraction Module.

Checks an uploaded resume and reads the text from it.
Supported formats: PDF and DOCX.
"""

import io
import logging
import os

from docx import Document
from pypdf import PdfReader

logging.getLogger("pypdf").setLevel(logging.ERROR)   # unreadable files are reported to the user instead

ALLOWED_TYPES = {".pdf": "PDF", ".docx": "DOCX"}
MAX_FILE_SIZE = 2 * 1024 * 1024   # 2 MB
MIN_TEXT_LENGTH = 50              # fewer characters than this means the file is unreadable


class ResumeError(Exception):
    """Raised with a message that can be shown directly to the user."""


def file_type_of(file_name):
    """Return 'PDF' or 'DOCX' for a supported file name, otherwise None."""
    extension = os.path.splitext(file_name)[1].lower()
    return ALLOWED_TYPES.get(extension)


def _text_from_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _text_from_docx(data):
    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    # Many resumes keep skills or education inside tables, so tables are read as well.
    for table in document.tables:
        for row in table.rows:
            lines.append(" , ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def extract_text(file_name, data):
    """
    Check the uploaded file and return (file type, extracted text).

    Raises ResumeError with a clear message if the file cannot be used.
    """
    file_type = file_type_of(file_name)
    if file_type is None:
        raise ResumeError("Please upload a supported PDF or DOCX file.")
    if len(data) == 0:
        raise ResumeError("The selected file is empty. Please choose another file.")
    if len(data) > MAX_FILE_SIZE:
        raise ResumeError("The file is larger than 2 MB. Please upload a smaller file.")

    try:
        text = _text_from_pdf(data) if file_type == "PDF" else _text_from_docx(data)
    except Exception:
        raise ResumeError("Text could not be extracted. Please try another file.")

    text = text.strip()
    if len(text) < MIN_TEXT_LENGTH:
        # Usually a scanned PDF: it holds a picture of the page, not text.
        raise ResumeError("Text could not be extracted. Please try another file.")
    return file_type, text
